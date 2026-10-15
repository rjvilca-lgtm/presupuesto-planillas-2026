"""
Plantilla única de Desarrollo Humano: docentes TC, sus horas y los salarios
2027 en un solo Excel, una fila por persona.

  - HORAS TC con valor: la persona es docente a tiempo completo (va al padrón
    HORAS_DOCENTES con esas horas). Vacía: no es docente TC.
  - Columnas de salario: van a Remuneraciones_2027.

El archivo es la lista completa: quien no esté se da de baja del padrón y de
salarios. La app no muestra datos: solo valida, informa errores e
inconsistencias (por número de fila y DNI) y cuenta altas, bajas y cambios.
La plantilla se descarga en blanco.
"""

import io
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .repositorio import COLS_REMUNERACION

# DNI (8 dígitos), carné de extranjería (9-12) o plaza nueva sin docente aún (NUEVO1...)
DOC_IDENTIDAD = re.compile(r"\d{8}|[0-9A-Za-z]{9,12}|NUEVO\d+", re.IGNORECASE)
HOJA = "Docentes"

ETIQUETAS_SALARIO = {
    "Basico": "BÁSICO", "Asignacion_Familiar": "ASIGNACIÓN FAMILIAR",
    "Bonificacion_Docente_TC": "BONIF. DOCENTE TC",
    "Bonificacion_Dedicacion_Exclusiva": "BONIF. DEDICACIÓN EXCLUSIVA",
    "Bono_por_Decano": "BONO POR DECANO",
    "Bono_por_Director_de_Escuela": "BONO POR DIRECTOR DE ESCUELA",
    "Bono_por_Departamento": "BONO POR DEPARTAMENTO", "Bono_al_Cargo": "BONO AL CARGO",
    "Bonificacion_Variable_por_Profesor_Extranjero": "BONIF. VARIABLE PROF. EXTRANJERO",
    "Bonificacion_Ordinaria_por_Acuerdo_Rem": "BONIF. ORDINARIA POR ACUERDO REM.",
    "Bonificacion_por_Representacion": "BONIF. POR REPRESENTACIÓN",
    "Grado_de_Maestria": "GRADO DE MAESTRÍA", "Grado_de_Doctorado": "GRADO DE DOCTORADO",
    "Asignacion_por_Vivienda": "ASIGNACIÓN POR VIVIENDA",
    "Asignacion_por_Hijo_Menor": "ASIGNACIÓN POR HIJO MENOR",
    "Vale_por_Alimentos": "VALE POR ALIMENTOS",
}
CABECERAS = {"DPTO": "DEPARTAMENTO", "DNI": "DNI", "NOMBRE": "DOCENTE",
             "CATEGORIA": "CATEGORÍA", "HORAS": "HORAS TC", **ETIQUETAS_SALARIO}


def _norm(texto) -> str:
    """Encabezado comparable: sin tildes, mayúsculas, solo letras y números."""
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", t.upper())


def normalizar_dni(valor) -> str:
    """Excel convierte el DNI en número y quita los ceros a la izquierda: se
    recuperan solo para 6-7 cifras (lo que Excel suele quitar)."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    t = str(int(valor)) if isinstance(valor, (int, float)) else str(valor).strip()
    if t.upper().startswith("NUEVO"):
        return t.upper()
    return t.zfill(8) if t.isdigit() and 6 <= len(t) < 8 else t


def unificar_dni(df: pd.DataFrame) -> pd.DataFrame:
    """DNI normalizado y sin duplicados (Remuneraciones tiene docentes con y sin
    el cero inicial; se queda la versión de 8 dígitos, la que usan
    HORAS_DOCENTES y los planes TC)."""
    df = df.copy()
    largo = df["DNI"].astype(str).str.strip().str.len()
    df["DNI"] = df["DNI"].map(normalizar_dni)
    return (df.assign(_l=largo).sort_values("_l", ascending=False)
              .drop_duplicates("DNI").drop(columns="_l").sort_index())


# ------------------------------------------------------------ plantilla

def plantilla(codigos: list[str]) -> bytes:
    """Plantilla en blanco con validaciones y una hoja de instrucciones."""
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    ws.append(list(CABECERAS.values()))
    n = len(CABECERAS)
    anchos = [16, 14, 42, 18, 10]
    for i, celda in enumerate(ws[1], start=1):
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="17233B" if i <= 5 else "7A1F2B")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = anchos[i - 1] if i <= 5 else 16
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "D2"
    for (celda,) in ws.iter_rows(min_row=2, max_row=1000, min_col=2, max_col=2):
        celda.number_format = "@"          # DNI como texto: conserva los ceros

    listas = wb.create_sheet("Listas")
    for i, c in enumerate(codigos, start=1):
        listas.cell(row=i, column=1, value=c)
    listas.sheet_state = "hidden"
    dv_dpto = DataValidation(type="list", formula1=f"=Listas!$A$1:$A${len(codigos)}",
                             allow_blank=True, showErrorMessage=True,
                             errorTitle="Departamento", error="Elige un departamento de la lista.")
    dv_dpto.add("A2:A1000")
    dv_horas = DataValidation(type="decimal", operator="between", formula1="1", formula2="40",
                              allow_blank=True, showErrorMessage=True, errorTitle="Horas TC",
                              error="Entre 1 y 40, o vacío si no es docente TC.")
    dv_horas.add("E2:E1000")
    dv_monto = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                              allow_blank=True, showErrorMessage=True, errorTitle="Monto",
                              error="Monto mayor o igual a 0, sin texto.")
    dv_monto.add(f"F2:{get_column_letter(n)}1000")
    for dv in (dv_dpto, dv_horas, dv_monto):
        ws.add_data_validation(dv)

    ins = wb.create_sheet("Instrucciones")
    ins.column_dimensions["A"].width = 115
    for i, linea in enumerate([
        "Docentes, horas y salarios 2027 — CONFIDENCIAL",
        "1. Una fila por persona. Este archivo es la LISTA COMPLETA: quien no esté aquí se "
        "da de baja del padrón de docentes TC y de salarios 2027 al importarlo.",
        "2. DEPARTAMENTO: código de la lista (DAIC, DHUM, ...), obligatorio para docentes TC. "
        "DNI: 8 dígitos, o 9 a 12 caracteres si es carné de extranjería; para una plaza "
        "nueva sin docente asignado usa NUEVO1, NUEVO2, ...",
        "3. HORAS TC: horas cronológicas del plan de trabajo (40, o menos si tiene carga "
        "parcial). Déjala vacía si la persona no es docente a tiempo completo.",
        "4. Columnas de salario (en rojo): montos MENSUALES en soles. Lo que no aplica va en 0 "
        "o vacío. Todo docente TC debe tener salario.",
        "5. No cambies los nombres de las columnas ni el nombre de la hoja «Docentes».",
        "6. Importa el archivo en la app: Desarrollo Humano › Importar. La app no muestra los "
        "datos: solo indica errores e inconsistencias por número de fila.",
        "7. No envíes este archivo por correo ni lo dejes en carpetas compartidas; bórralo "
        "después de importarlo.",
    ], start=1):
        c = ins.cell(row=i, column=1, value=linea)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=13)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------ lectura

def filas_txt(indices) -> str:
    indices = sorted(set(indices))
    return ", ".join(map(str, indices[:15])) + (f" y {len(indices) - 15} más"
                                                if len(indices) > 15 else "")


def leer(archivo, codigos: set) -> tuple[pd.DataFrame | None, list[str]]:
    """(datos, errores bloqueantes). El índice del DataFrame es el número de
    fila del Excel, para poder informarlo."""
    try:
        libro = pd.read_excel(archivo, sheet_name=None, dtype=object)
    except Exception as ex:
        return None, [f"No se pudo leer el archivo como Excel ({type(ex).__name__})."]
    df = libro.get(HOJA)
    if df is None:
        df = next(iter(libro.values()))
    esperadas = {_norm(v): k for k, v in CABECERAS.items()}
    df = df.rename(columns={c: esperadas.get(_norm(c), c) for c in df.columns})
    faltan = [CABECERAS[k] for k in CABECERAS if k not in df.columns]
    if faltan:
        return None, ["Faltan columnas: " + ", ".join(faltan)
                      + ". Usa la plantilla sin cambiar los encabezados."]
    df = df[list(CABECERAS)].dropna(how="all")
    df.index = df.index + 2   # la fila 1 del Excel es el encabezado
    if df.empty:
        return None, ["El archivo no tiene filas."]

    errores = []
    df["DNI"] = df["DNI"].map(normalizar_dni)
    df["DPTO"] = df["DPTO"].map(lambda v: str(v).strip().upper() if pd.notna(v) else "")
    for c in ("NOMBRE", "CATEGORIA"):
        df[c] = df[c].map(lambda v: " ".join(str(v).split()) if pd.notna(v) else "")

    horas_txt = df["HORAS"]
    df["HORAS"] = pd.to_numeric(horas_txt, errors="coerce")
    texto = list(df.index[horas_txt.notna() & df["HORAS"].isna()])
    for c in COLS_REMUNERACION:
        num = pd.to_numeric(df[c], errors="coerce")
        texto += list(df.index[df[c].notna() & num.isna()])
        df[c] = num.fillna(0.0).round(2)
    df["TOTAL"] = df[COLS_REMUNERACION].sum(axis=1)
    tc = df["HORAS"].notna()

    def error(mascara, mensaje):
        filas = df.index[mascara]
        if len(filas):
            errores.append(f"{mensaje}: filas {filas_txt(filas)}.")

    if texto:
        errores.append(f"Valores con texto donde va un número: filas {filas_txt(texto)}.")
    error(df["DNI"] == "", "Sin DNI")
    error((df["DNI"] != "") & ~df["DNI"].map(lambda d: bool(DOC_IDENTIDAD.fullmatch(d))),
          "DNI inválido")
    error(df["DNI"].duplicated(keep=False) & (df["DNI"] != ""), "DNI repetido")
    error(df["NOMBRE"] == "", "Sin nombre del docente")
    error((df["DPTO"] != "") & ~df["DPTO"].isin(codigos), "Departamento que no existe")
    error(tc & (df["DPTO"] == ""), "Docente TC sin departamento")
    error(tc & ((df["HORAS"] < 1) | (df["HORAS"] > 40)), "Horas TC fuera de 1 a 40")
    error((df[COLS_REMUNERACION] < 0).any(axis=1), "Montos negativos")
    error(~tc & (df["TOTAL"] <= 0), "Fila sin horas TC ni salario")
    # La inconsistencia que más importa: sin salario, su sueldo TC se valoriza en 0
    error(tc & (df["TOTAL"] <= 0), "Docente TC sin salario (su sueldo se valorizaría en 0)")
    return (None if errores else df), errores


def contar_movimientos(actual: pd.DataFrame, nuevo: pd.DataFrame, campos: list[str]) -> dict:
    """Altas, bajas y cambios por DNI (solo conteos: la app no muestra datos)."""
    a = actual.drop_duplicates("DNI").set_index("DNI")
    n = nuevo.drop_duplicates("DNI").set_index("DNI")
    cambios = 0
    for dni in set(a.index) & set(n.index):
        for c in campos:
            va, vn = a.at[dni, c], n.at[dni, c]
            if isinstance(va, (int, float)) or isinstance(vn, (int, float)):
                va = 0.0 if va is None or pd.isna(va) else float(va)
                vn = 0.0 if vn is None or pd.isna(vn) else float(vn)
                distinto = abs(va - vn) > 0.005
            else:
                distinto = " ".join(str(va or "").split()) != " ".join(str(vn or "").split())
            if distinto:
                cambios += 1
                break
    return {"altas": len(set(n.index) - set(a.index)),
            "bajas": len(set(a.index) - set(n.index)), "cambios": cambios}
