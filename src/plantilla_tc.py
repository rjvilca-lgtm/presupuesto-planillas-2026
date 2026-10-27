"""
Plan de trabajo TC en Excel: descargar la ficha de un semestre, llenarla fuera
de la app y volver a cargarla.

La plantilla sale de la ficha tal como está en pantalla (padrón del
departamento × estructura CECO -> concepto, con las horas ya registradas):
una fila por docente, centro de costo y concepto. Solo se llena la columna
HORAS, en hh:mm. Las horas cronológicas de dictado no van: se calculan de las
académicas (× 45/60), igual que en la ficha. La hoja Resumen suma lo de cada
docente con fórmulas para ver, sin la app, a quién le faltan horas.

Al cargarla, la app solo pasa las horas a la ficha (no guarda): se revisan y
se pulsa Guardar, con las mismas validaciones de siempre. Si el archivo tiene
un error, no se carga nada.
"""

import io
from datetime import datetime, time, timedelta

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from .config import (CONCEPTO_ACADEMICAS, CONCEPTO_CRONOLOGICAS, FACTOR_CRONOLOGICAS,
                     cfg_de, de_hhmm, es_concepto)
from .plantillas import _norm, normalizar_dni

HOJA = "Plan TC"
CABECERAS = ["DNI", "DOCENTE", "CENTRO DE COSTO", "DESCRIPCIÓN DEL CENTRO DE COSTO",
             "CONCEPTO", "HORAS (hh:mm)"]
FILA_CABECERA = 4
MAX_HORAS = 48          # por concepto: más es casi seguro un error de tipeo
_FORMATO_HORAS = "[hh]:mm"
_AZUL, _EDITABLE = "17233B", "FFF7E0"


def plantilla(ledger: pd.DataFrame, hoja: str, origen: str, desc_ceco: dict) -> bytes:
    """Excel de la ficha de un semestre y departamento, con las horas actuales."""
    cfg = cfg_de(hoja)
    filas = ledger[~ledger["CONCEPTO"].map(lambda c: es_concepto(c, CONCEPTO_CRONOLOGICAS))]
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    ws["A1"] = f"{cfg['etiqueta']} · {origen}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Llena solo la columna HORAS (amarilla) en hh:mm, por ejemplo 03:30. "
                "Las horas académicas de dictado se convierten solas a cronológicas "
                "(× 45/60). Revisa la hoja Resumen antes de cargar el archivo en la app.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 32
    for i, (titulo, ancho) in enumerate(zip(CABECERAS, [12, 38, 18, 44, 40, 14]), start=1):
        c = ws.cell(row=FILA_CABECERA, column=i, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = ancho
    ws.row_dimensions[FILA_CABECERA].height = 30
    ws.freeze_panes = ws.cell(row=FILA_CABECERA + 1, column=3)

    r = FILA_CABECERA
    for r, fila in enumerate(filas.itertuples(index=False), start=FILA_CABECERA + 1):
        ws.cell(row=r, column=1, value=str(fila.DNI)).number_format = "@"
        ws.cell(row=r, column=2, value=fila.NOMBRE)
        ws.cell(row=r, column=3, value=str(fila.CENTRO_COSTO)).number_format = "@"
        ws.cell(row=r, column=4, value=desc_ceco.get(fila.CENTRO_COSTO, fila.CENTRO_COSTO))
        ws.cell(row=r, column=5, value=fila.CONCEPTO)
        h = ws.cell(row=r, column=6, value=float(fila.HORAS or 0) / 24)
        h.number_format = _FORMATO_HORAS
        h.fill = PatternFill("solid", fgColor=_EDITABLE)
        h.protection = Protection(locked=False)
    ultima = r
    if ultima > FILA_CABECERA:
        dv = DataValidation(type="decimal", operator="between", formula1="0",
                            formula2=str(MAX_HORAS / 24), allow_blank=True,
                            showErrorMessage=True, errorTitle="Horas",
                            error=f"Escribe las horas como hh:mm, por ejemplo 03:30 "
                                  f"(hasta {MAX_HORAS}:00).")
        dv.add(f"F{FILA_CABECERA + 1}:F{ultima}")
        ws.add_data_validation(dv)
        ws.auto_filter.ref = f"A{FILA_CABECERA}:F{ultima}"
    # Protegida sin contraseña: evita cambiar por error DNI, centros de costo o
    # conceptos (la app no los reconocería); se puede quitar desde Revisar.
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False

    _resumen(wb, ledger, ultima)

    # Para comprobar al cargar que el archivo es de esta ficha
    app = wb.create_sheet("_app")
    app["A1"], app["A2"] = hoja, origen
    app.sheet_state = "hidden"

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _resumen(wb, ledger: pd.DataFrame, ultima: int):
    ws = wb.create_sheet("Resumen")
    titulos = ["DNI", "DOCENTE", "HORAS QUE DEBE SUMAR", "HORAS ASIGNADAS",
               "DIFERENCIA (h)", "ESTADO"]
    for i, (t, ancho) in enumerate(zip(titulos, [12, 38, 14, 14, 14, 14]), start=1):
        c = ws.cell(row=1, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = ancho
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    rango = lambda col: f"'{HOJA}'!${col}${FILA_CABECERA + 1}:${col}${max(ultima, FILA_CABECERA + 1)}"
    horas, dni, concepto = rango("F"), rango("A"), rango("E")
    docentes = ledger.drop_duplicates("DNI")[["DNI", "NOMBRE", "HORAS_META"]]
    for r, d in enumerate(docentes.itertuples(index=False), start=2):
        ws.cell(row=r, column=1, value=str(d.DNI)).number_format = "@"
        ws.cell(row=r, column=2, value=d.NOMBRE)
        ws.cell(row=r, column=3, value=float(d.HORAS_META) / 24).number_format = _FORMATO_HORAS
        ws.cell(row=r, column=4, value=(
            f'=SUMIFS({horas},{dni},A{r},{concepto},"<>{CONCEPTO_ACADEMICAS}")'
            f'+SUMIFS({horas},{dni},A{r},{concepto},"{CONCEPTO_ACADEMICAS}")'
            f"*{FACTOR_CRONOLOGICAS}")).number_format = _FORMATO_HORAS
        ws.cell(row=r, column=5, value=f"=ROUND((C{r}-D{r})*24,2)").number_format = "0.00"
        ws.cell(row=r, column=6,
                value=f'=IF(ABS(E{r})<=0.01,"Completo",IF(E{r}>0,"Faltan","Sobran"))')
    ws.cell(row=len(docentes) + 3, column=1,
            value="Diferencia: horas que faltan (positivo) o sobran (negativo), en horas "
                  "decimales. Las académicas cuentan × 45/60.").font = Font(italic=True)


def _horas(valor) -> float | None:
    """Lo que Excel devuelve en una celda de horas -> horas en decimal.
    None si no se entiende."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return 0.0
    if isinstance(valor, timedelta):
        return valor.total_seconds() / 3600
    if isinstance(valor, datetime):          # más de 24 h: fecha desde 1899-12-30
        return (valor - datetime(1899, 12, 30)).total_seconds() / 3600
    if isinstance(valor, time):
        return valor.hour + valor.minute / 60 + valor.second / 3600
    if isinstance(valor, (int, float)):
        return float(valor) * 24              # celda con formato de hora: días
    return de_hhmm(valor)


def leer(archivo, ledger: pd.DataFrame, hoja: str, origen: str):
    """Pasa las horas del Excel a una copia de la ficha (ledger).
    Devuelve (ledger nuevo o None, errores, resumen). Con errores no cambia
    nada. Lo que no esté en el archivo queda como estaba en la ficha."""
    try:
        wb = load_workbook(archivo, data_only=True)
    except Exception as ex:
        return None, [f"No se pudo abrir el archivo como Excel (.xlsx): {ex}"], ""
    if HOJA not in wb.sheetnames:
        return None, [f"El archivo no tiene la hoja «{HOJA}». Descarga la plantilla desde "
                      "esta pantalla y llénala."], ""
    if "_app" in wb.sheetnames:
        h, o = wb["_app"]["A1"].value, wb["_app"]["A2"].value
        if (h, o) != (hoja, origen):
            otra = cfg_de(h)["etiqueta"] if h in ("1. Plan TC-I", "Plan TC-II") else h
            return None, [f"El archivo es de «{otra} · {o}» y estás en "
                          f"«{cfg_de(hoja)['etiqueta']} · {origen}». Cárgalo en su pestaña."], ""
    ws = wb[HOJA]
    cab = [_norm(c.value) for c in ws[FILA_CABECERA]]
    try:
        i_dni, i_ceco, i_con = (cab.index(_norm(t)) for t in
                                ("DNI", "CENTRO DE COSTO", "CONCEPTO"))
        i_h = next(i for i, c in enumerate(cab) if c.startswith("HORAS"))
    except (ValueError, StopIteration):
        return None, [f"No se encontraron las columnas DNI, CENTRO DE COSTO, CONCEPTO y "
                      f"HORAS en la fila {FILA_CABECERA}. No cambies los encabezados."], ""

    nuevo = ledger.copy()
    indice = {(str(d), str(c), str(k).strip().lower()): i for i, d, c, k in zip(
        nuevo.index, nuevo["DNI"], nuevo["CENTRO_COSTO"], nuevo["CONCEPTO"])}
    errores, vistos, cambios = [], set(), 0
    for r, fila in enumerate(ws.iter_rows(min_row=FILA_CABECERA + 1, values_only=True),
                             start=FILA_CABECERA + 1):
        if fila is None or all(v is None or str(v).strip() == "" for v in fila):
            continue
        dni = normalizar_dni(fila[i_dni])
        ceco = str(fila[i_ceco] or "").strip()
        concepto = str(fila[i_con] or "").strip()
        if es_concepto(concepto, CONCEPTO_CRONOLOGICAS):
            continue                          # se calcula de las académicas
        clave = (dni, ceco, concepto.lower())
        if clave not in indice:
            errores.append(f"Fila {r}: el docente {dni or '(sin DNI)'} con el centro de costo "
                           f"{ceco or '(vacío)'} y el concepto «{concepto}» no está en la "
                           "ficha. No agregues ni cambies filas: solo llena las horas.")
            continue
        if clave in vistos:
            errores.append(f"Fila {r}: el docente {dni}, {ceco} · {concepto} está repetido.")
            continue
        vistos.add(clave)
        valor = _horas(fila[i_h])
        if valor is None or valor < 0 or valor > MAX_HORAS:
            errores.append(f"Fila {r}: «{fila[i_h]}» no es una hora válida para el docente "
                           f"{dni}. Usa hh:mm, por ejemplo 03:30.")
            continue
        valor = round(valor * 60) / 60        # a minutos, como la ficha
        i = indice[clave]
        if abs(float(nuevo.at[i, "HORAS"] or 0) - valor) > 1e-9:
            cambios += 1
        nuevo.at[i, "HORAS"] = valor
    if errores:
        return None, errores, ""
    if not vistos:
        return None, ["El archivo no tiene filas con docentes."], ""

    # Cronológicas = académicas × 45/60, por docente y centro de costo
    es_acad = nuevo["CONCEPTO"].map(lambda c: es_concepto(c, CONCEPTO_ACADEMICAS))
    es_cron = nuevo["CONCEPTO"].map(lambda c: es_concepto(c, CONCEPTO_CRONOLOGICAS))
    acad = nuevo[es_acad].set_index(["DNI", "CENTRO_COSTO"])["HORAS"]
    claves = list(zip(nuevo.loc[es_cron, "DNI"], nuevo.loc[es_cron, "CENTRO_COSTO"]))
    nuevo.loc[es_cron, "HORAS"] = [float(acad.get(k, 0.0)) * FACTOR_CRONOLOGICAS
                                   for k in claves]
    docentes = len({d for d, _, _ in vistos})
    return nuevo, [], (f"{len(vistos)} fila(s) de {docentes} docente(s); "
                       f"{cambios} hora(s) distintas a las de la ficha.")

