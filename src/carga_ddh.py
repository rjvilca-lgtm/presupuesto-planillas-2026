"""
Valorización de Desarrollo Humano (DDH) 2027: plantilla, lectura y validación.

Reemplaza la macro SubirDatos_DDH de "Subir datos DDH.xlsm". Mismas columnas
que esa hoja (CENTRO_COSTO, PARTIDA, DIMENSION, MES, IMPORTE, CONCEPTO): se
puede pegar lo que DH ya prepara. Al importar se reemplaza toda la
valorización DDH 2027 en PLANILLA (sql/crear_carga_ddh_2027.sql).

CONCEPTO puede llevar nombres de personas: la app no muestra filas, solo
errores por número de fila y totales agregados.
"""

import io
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .plantillas import filas_txt

HOJA = "Datos"
COLUMNAS = ["CENTRO_COSTO", "PARTIDA", "DIMENSION", "MES", "IMPORTE", "CONCEPTO"]
ANCHOS = [18, 16, 12, 8, 14, 50]
LARGO_MAX = {"CENTRO_COSTO": 50, "PARTIDA": 50, "CONCEPTO": 100}
IMPORTE_MAX = 99_999_999.99   # PLANILLA.IMPORTE es decimal(10,2)


def _norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", t.upper())


def _texto(v) -> str:
    """Celda -> texto limpio; los números enteros sin '.0' (DIMENSION 20, no 20.0)."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return " ".join(str(v).split())


# ------------------------------------------------------------ plantilla

def plantilla(dimensiones: pd.DataFrame) -> bytes:
    """dimensiones: CODIGO (el que va en la columna) y DESCRIPCION."""
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    ws.append(COLUMNAS)
    for i, (celda, ancho) in enumerate(zip(ws[1], ANCHOS), start=1):
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="17233B")
        celda.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[celda.column_letter].width = ancho
    ws.freeze_panes = "A2"
    for fila in ws.iter_rows(min_row=2, max_row=20000, min_col=1, max_col=2):
        for celda in fila:
            celda.number_format = "@"   # códigos como texto

    dim = wb.create_sheet("Dimensiones")
    dim.append(["DIMENSION", "DESCRIPCIÓN"])
    for r in dimensiones.itertuples(index=False):
        dim.append([str(r.CODIGO), str(r.DESCRIPCION)])
    dim.column_dimensions["A"].width = 12
    dim.column_dimensions["B"].width = 45
    for celda in dim[1]:
        celda.font = Font(bold=True)
    n = len(dimensiones) + 1
    validaciones = [
        (DataValidation(type="list", formula1=f"=Dimensiones!$A$2:$A${n}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Dimensión",
                        error="Elige una dimensión de la hoja Dimensiones."), "C2:C20000"),
        (DataValidation(type="whole", operator="between", formula1="1", formula2="12",
                        allow_blank=True, showErrorMessage=True, errorTitle="Mes",
                        error="Mes de 1 a 12."), "D2:D20000"),
        (DataValidation(type="decimal", operator="between", formula1="-99999999",
                        formula2="99999999", allow_blank=True, showErrorMessage=True,
                        errorTitle="Importe", error="Importe en soles, sin texto."), "E2:E20000"),
    ]
    for dv, rango in validaciones:
        dv.add(rango)
        ws.add_data_validation(dv)

    ins = wb.create_sheet("Instrucciones")
    ins.column_dimensions["A"].width = 115
    for i, linea in enumerate([
        "Valorización de Desarrollo Humano 2027 — CONFIDENCIAL",
        "1. Una fila por centro de costo, partida, dimensión, mes y concepto, con el importe "
        "en soles. Son las mismas columnas de la hoja «Subir datos DDH»: puedes pegar ahí lo "
        "que ya preparas.",
        "2. CENTRO_COSTO: código 2027 (01.xx.xx.xx.xx). Si pegas un código antiguo (03.xx) que "
        "tenga equivalencia, la app lo convierte solo y te avisa cuántos convirtió.",
        "3. PARTIDA: código de partida presupuestal (94.1.1.1.xxx). DIMENSION: código de la "
        "hoja «Dimensiones» (20 = Remuneración Docencia, 22 = Administrativa, ...).",
        "4. MES: número de 1 a 12. IMPORTE: soles, con hasta 2 decimales. CONCEPTO: texto "
        "libre de hasta 100 caracteres (opcional).",
        "5. Este archivo es la valorización COMPLETA de DDH 2027: al importarlo reemplaza "
        "la anterior.",
        "6. Importa el archivo en la app: Desarrollo Humano › Valorización DDH. La app no "
        "muestra las filas: solo errores por número de fila y totales.",
        "7. No envíes este archivo por correo ni lo dejes en carpetas compartidas; bórralo "
        "después de importarlo.",
    ], start=1):
        c = ins.cell(row=i, column=1, value=linea)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=13)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------ lectura

def leer(archivo, cecos_2027: set, equivalencia: dict, dimensiones: set,
         partidas: set) -> tuple[pd.DataFrame | None, list[str], list[str]]:
    """(datos, errores que impiden importar, avisos). El índice del DataFrame
    es el número de fila del Excel."""
    try:
        libro = pd.read_excel(archivo, sheet_name=None, dtype=object)
    except Exception as ex:
        return None, [f"No se pudo leer el archivo como Excel ({type(ex).__name__})."], []
    # La hoja «Datos» de la plantilla, o la primera que tenga las columnas
    # (la hoja «Hoja1» del Excel antiguo también sirve)
    df = None
    for nombre, hoja in sorted(libro.items(), key=lambda kv: kv[0] != HOJA):
        cols = {_norm(c): c for c in hoja.columns}
        if all(_norm(c) in cols for c in COLUMNAS):
            df = hoja.rename(columns={cols[_norm(c)]: c for c in COLUMNAS})[COLUMNAS]
            break
    if df is None:
        return None, ["No se encontró una hoja con las columnas " + ", ".join(COLUMNAS)
                      + ". Usa la plantilla sin cambiar los encabezados."], []
    df = df.dropna(how="all")
    df.index = df.index + 2
    if df.empty:
        return None, ["El archivo no tiene filas."], []

    errores, avisos = [], []

    def error(mascara, mensaje):
        filas = df.index[mascara]
        if len(filas):
            errores.append(f"{mensaje}: filas {filas_txt(filas)}.")

    for c in ("CENTRO_COSTO", "PARTIDA", "DIMENSION", "CONCEPTO"):
        df[c] = df[c].map(_texto)
    mes = pd.to_numeric(df["MES"], errors="coerce")
    importe = pd.to_numeric(df["IMPORTE"], errors="coerce")

    # Códigos antiguos con equivalencia: se convierten al código 2027
    antiguo = ~df["CENTRO_COSTO"].isin(cecos_2027) & df["CENTRO_COSTO"].isin(equivalencia)
    if antiguo.any():
        avisos.append(f"{int(antiguo.sum())} fila(s) con {df.loc[antiguo, 'CENTRO_COSTO'].nunique()} "
                      "centro(s) de costo de la codificación antigua se convirtieron al código "
                      "2027 según la tabla de equivalencias.")
        df.loc[antiguo, "CENTRO_COSTO"] = df.loc[antiguo, "CENTRO_COSTO"].map(equivalencia)

    error(df["CENTRO_COSTO"] == "", "Sin centro de costo")
    sin_cc = (df["CENTRO_COSTO"] != "") & ~df["CENTRO_COSTO"].isin(cecos_2027)
    if sin_cc.any():
        codigos = sorted(df.loc[sin_cc, "CENTRO_COSTO"].unique())
        errores.append("Centros de costo que no existen en 2027 ni tienen equivalencia: "
                       + ", ".join(codigos[:10]) + (" …" if len(codigos) > 10 else "")
                       + f" (filas {filas_txt(df.index[sin_cc])}).")
    error(df["PARTIDA"] == "", "Sin partida")
    error((df["DIMENSION"] != "") & ~df["DIMENSION"].isin(dimensiones),
          "Dimensión que no está en la lista")
    error(df["DIMENSION"] == "", "Sin dimensión")
    error(df["MES"].notna() & mes.isna(), "Mes con texto")
    error(df["MES"].isna() | ((mes.notna()) & (~mes.between(1, 12) | (mes % 1 != 0))),
          "Mes vacío o fuera de 1 a 12")
    error(df["IMPORTE"].notna() & importe.isna(), "Importe con texto")
    error(df["IMPORTE"].isna(), "Sin importe")
    error(importe.abs() > IMPORTE_MAX, "Importe demasiado grande")
    for c, largo in LARGO_MAX.items():
        error(df[c].str.len() > largo, f"{c} con más de {largo} caracteres")

    desconocidas = sorted(set(df["PARTIDA"]) - partidas - {""})
    if desconocidas:
        avisos.append("Partidas que no están en el catálogo de partidas presupuestales "
                      "(se importan igual; revisa que sean correctas): " + ", ".join(desconocidas))
    negativos = importe < 0
    if negativos.any():
        avisos.append(f"{int(negativos.sum())} fila(s) con importe negativo (filas "
                      f"{filas_txt(df.index[negativos])}).")
    ceros = importe == 0
    if ceros.any():
        avisos.append(f"{int(ceros.sum())} fila(s) con importe 0: no se importan.")

    if errores:
        return None, errores, avisos
    df["MES"] = mes.astype(int)
    df["IMPORTE"] = importe.round(2)
    df = df[df["IMPORTE"] != 0]
    if df.empty:
        return None, ["Todas las filas tienen importe 0."], avisos
    return df, [], avisos
